#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sinh file HD_DICH_THUAT_MOCK.docx - bo hop dong bien dich mock tron ven.
Hoc theo hop dong that HDDG_2025-b_i_gia_bin-__t_kho.docx + bo sung cac dieu
khoan bao mat / so huu tri tue / cham dut / tranh chap theo mau hop dong that khac.
Placeholder dung cu phap docxtemplater mac dinh: {ten_bien} (ngoac don).
Khong dung thong tin ca nhan that - ben B chi de placeholder.
Chay: python3 build_mock_contract.py  (ghi de HD_DICH_THUAT_MOCK.docx)
"""
import os
from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "HD_DICH_THUAT_MOCK.docx")

d = Document()

# Kho giay A4 (khop @page size A4 khi in tu docx-preview, margin 0)
sec = d.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)

# Font chu mac dinh
style = d.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(13)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15


def add(text, bold=False, center=False, italic=False, size=None, space_after=None):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text):
    return add(text, bold=True, size=13, space_after=6)


def body(text):
    return add(text)


def bullet(text, level=0):
    p = d.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.left_indent = Cm(1.0 + level * 0.75)
    return p


# ---------- Tieu de ----------
add("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, center=True, size=13, space_after=0)
add("Độc lập - Tự do - Hạnh phúc", center=True, size=13, space_after=12)
add("HỢP ĐỒNG BIÊN DỊCH ẤN PHẨM", bold=True, center=True, size=15, space_after=0)
add("Số: {so_hop_dong}", center=True, size=13, space_after=12)

# ---------- Can cu ----------
body("Căn cứ Bộ luật Dân sự số 91/2015/QH13 ngày 24/11/2015 của Quốc hội nước Cộng hòa xã hội chủ nghĩa Việt Nam;")
body("Căn cứ Luật Sở hữu trí tuệ số 50/2005/QH11 ngày 29/11/2005 của Quốc hội nước Cộng hòa xã hội chủ nghĩa Việt Nam;")
body("Căn cứ Luật Xuất bản số 19/2012/QH13 ngày 20/11/2012 của Quốc hội nước Cộng hòa xã hội chủ nghĩa Việt Nam;")
body("Căn cứ Nghị định số 195/2013/NĐ-CP ngày 21/11/2013 của Chính phủ quy định chi tiết một số điều và biện pháp thi hành Luật Xuất bản;")
body("Căn cứ vào khả năng, nhu cầu và sự thỏa thuận của hai bên.")
add("")

add("Hôm nay, ngày {ngay_ky_ngay} tháng {ngay_ky_thang} năm {ngay_ky_nam}, tại {dia_diem_ky}, hai bên chúng tôi gồm:", space_after=8)

# ---------- Ben A ----------
add("BÊN A: {ben_a_ten}", bold=True, space_after=4)
bullet("Đại diện: {ben_a_dai_dien} - Chức vụ: {ben_a_chuc_vu}")
bullet("Địa chỉ: {ben_a_dia_chi}")
bullet("Điện thoại: {ben_a_sdt} - Email: {ben_a_email}")
bullet("Mã số thuế: {ben_a_mst}")
add("")

# ---------- Ben B ----------
add("BÊN B: DỊCH GIẢ", bold=True, space_after=4)
bullet("Họ và tên: {ben_b_ten} - Bút danh: {ben_b_but_danh}")
bullet("Số CCCD: {ben_b_cccd} (Ngày cấp: {ben_b_cccd_ngay_cap} - Nơi cấp: {ben_b_cccd_noi_cap})")
bullet("Địa chỉ thường trú: {ben_b_dia_chi}")
bullet("Điện thoại: {ben_b_sdt} - Email: {ben_b_email}")
bullet("Số tài khoản: {ben_b_stk} - Ngân hàng: {ben_b_ngan_hang}")
add("")

body('Sau khi bàn bạc, trao đổi, hai bên thống nhất ký kết Hợp đồng biên dịch ấn phẩm "{ten_tac_pham}" của tác giả {tac_gia} (nguyên tác tiếng {ngon_ngu_goc}) với các điều khoản sau đây:')
add("")

# ---------- Dieu 1 ----------
heading("ĐIỀU 1. NỘI DUNG CÔNG VIỆC")
body('1.1. Bên B đồng ý thực hiện biên dịch toàn bộ ấn phẩm "{ten_tac_pham}" của tác giả {tac_gia} từ tiếng {ngon_ngu_goc} sang tiếng Việt cho Bên A, bao gồm toàn bộ nội dung chính văn, lời tựa, lời bạt và các phần phụ lục kèm theo (nếu có).')
body("1.2. Bản dịch phải bảo đảm đầy đủ, chính xác, trung thành với nội dung và tinh thần của nguyên tác; văn phong tiếng Việt trong sáng, phù hợp với đối tượng bạn đọc và thuần phong mỹ tục Việt Nam.")
body("1.3. Bên B chịu trách nhiệm về tính pháp lý của nội dung bản dịch; không được tự ý cắt xén, thêm bớt làm sai lệch nội dung nguyên tác nếu không có sự đồng ý bằng văn bản của Bên A.")
add("")

# ---------- Dieu 2 ----------
heading("ĐIỀU 2. THỜI HẠN THỰC HIỆN")
body("2.1. Bên B hoàn thành và bàn giao toàn bộ bản dịch cho Bên A chậm nhất vào ngày {ngay_giao_ban_dich}, tức trong vòng {thoi_han_ngay} ngày kể từ ngày ký Hợp đồng này.")
body("2.2. Trong thời hạn {thoi_han_nghiem_thu} ngày kể từ ngày nhận bản dịch, Bên A thông báo kết quả nghiệm thu cho Bên B. Trường hợp bản dịch cần chỉnh sửa, Bên B có trách nhiệm hoàn thiện trong vòng {thoi_han_chinh_sua} ngày kể từ ngày nhận yêu cầu bằng văn bản của Bên A.")
body("2.3. Trường hợp bất khả kháng hoặc có lý do chính đáng được hai bên chấp thuận bằng văn bản, thời hạn nêu tại khoản 2.1 và 2.2 được gia hạn tương ứng.")
add("")

# ---------- Dieu 3 ----------
heading("ĐIỀU 3. GIÁ TRỊ HỢP ĐỒNG, THUẾ VÀ PHƯƠNG THỨC THANH TOÁN")
body("3.1. Giá trị Hợp đồng: {gia_tri_hd_so} đồng (Bằng chữ: {gia_tri_hd_chu}).")
body("3.2. Thuế thu nhập cá nhân: Giá trị Hợp đồng đã bao gồm thuế thu nhập cá nhân của Bên B. Bên A có trách nhiệm khấu trừ 10% thuế thu nhập cá nhân tại nguồn, kê khai và nộp thay cho Bên B theo quy định của pháp luật. Số tiền Bên B thực nhận sau khi khấu trừ thuế là {thuc_nhan_so} đồng (Bằng chữ: {thuc_nhan_chu}).")
body("3.3. Thanh toán: Bên A thanh toán 100% giá trị Hợp đồng (sau khi đã khấu trừ thuế theo khoản 3.2) cho Bên B trong vòng 30 ngày kể từ ngày có quyết định phát hành ấn phẩm, hoặc 04 tháng kể từ ngày Bên B bàn giao toàn bộ bản dịch và được nghiệm thu, tùy theo mốc nào đến trước.")
body("3.4. Phương thức thanh toán: chuyển khoản vào tài khoản ngân hàng của Bên B ghi tại phần thông tin các bên, hoặc bằng tiền mặt, bằng đồng Việt Nam.")
add("")

# ---------- Dieu 4 ----------
heading("ĐIỀU 4. QUYỀN VÀ NGHĨA VỤ CỦA BÊN A")
body("4.1. Nghĩa vụ của Bên A:")
bullet("Cung cấp cho Bên B đầy đủ bản thảo nguyên tác, tài liệu tham khảo và các thông tin cần thiết để thực hiện công việc;")
bullet("Thông báo kết quả nghiệm thu và yêu cầu chỉnh sửa (nếu có) trong thời hạn quy định tại Điều 2;")
bullet("Thanh toán đầy đủ, đúng thời hạn theo quy định tại Điều 3;")
bullet("Khi biên tập, hiệu đính bản dịch, thông báo và trao đổi thống nhất với Bên B những nội dung cần sửa chữa trên tinh thần xây dựng;")
bullet("Tuân thủ đầy đủ các quy định của Luật Xuất bản trong quá trình xuất bản, phát hành ấn phẩm.")
body("4.2. Quyền của Bên A:")
bullet("Được nhận bản dịch đầy đủ, đúng chất lượng và đúng tiến độ theo Hợp đồng;")
bullet("Từ chối nghiệm thu và chấm dứt Hợp đồng nếu bản dịch có sai sót nghiêm trọng không thể khắc phục;")
bullet("Sở hữu toàn bộ bản dịch và giữ bản quyền bản dịch theo quy định tại Điều 6.")
body("4.3. Hai bên xác nhận đây là hợp đồng khoán việc/dịch vụ. Bên A không có nghĩa vụ tham gia, đóng bảo hiểm xã hội, bảo hiểm y tế, bảo hiểm thất nghiệp hoặc các khoản bảo hiểm bắt buộc khác cho Bên B. Bên B tự thực hiện các nghĩa vụ liên quan theo quy định của pháp luật.")
add("")

# ---------- Dieu 5 ----------
heading("ĐIỀU 5. QUYỀN VÀ NGHĨA VỤ CỦA BÊN B")
body("5.1. Nghĩa vụ của Bên B:")
bullet("Thực hiện công việc đúng nội dung, chất lượng và tiến độ quy định tại Điều 1 và Điều 2;")
bullet("Bàn giao bản dịch dưới dạng file điện tử theo định dạng do Bên A yêu cầu; phối hợp chỉnh sửa theo kết quả nghiệm thu;")
bullet("Bảo mật tuyệt đối nội dung bản thảo và bản dịch; không được gửi một phần hoặc toàn bộ bản dịch cho bên thứ ba dưới bất kỳ hình thức nào nếu chưa được sự đồng ý bằng văn bản của Bên A;")
bullet("Chịu trách nhiệm về tính chính xác, hợp pháp của nội dung dịch thuật.")
body("5.2. Quyền của Bên B:")
bullet("Được đứng tên (hoặc bút danh {ben_b_but_danh}) trên ấn phẩm với tư cách dịch giả;")
bullet("Được nhận đầy đủ thù lao theo quy định tại Điều 3;")
bullet("Được Bên A cung cấp tài liệu, thông tin cần thiết để thực hiện công việc.")
add("")

# ---------- Dieu 6 ----------
heading("ĐIỀU 6. BẢO MẬT VÀ SỞ HỮU TRÍ TUỆ")
body("6.1. Toàn bộ bản thảo, bản dịch, tài liệu và thông tin mà Bên A cung cấp cho Bên B trong quá trình thực hiện Hợp đồng là tài liệu mật. Bên B cam kết không tiết lộ, sao chép, sử dụng cho mục đích khác ngoài phạm vi Hợp đồng này.")
body("6.2. Quyền sở hữu bản dịch, bao gồm quyền khai thác, xuất bản, tái bản, chuyển thể và các quyền tài sản khác phát sinh từ bản dịch, thuộc về Bên A. Bên A giữ bản quyền bản dịch vĩnh viễn.")
body("6.3. Bên B không được sử dụng lại, sao chép, công bố, chuyển nhượng hoặc khai thác bản dịch dưới bất kỳ hình thức nào nếu không có sự đồng ý bằng văn bản của Bên A.")
body("6.4. Nghĩa vụ bảo mật tại Điều này vẫn có hiệu lực sau khi Hợp đồng chấm dứt hoặc hết hiệu lực.")
add("")

# ---------- Dieu 7 ----------
heading("ĐIỀU 7. CHẤM DỨT HỢP ĐỒNG VÀ GIẢI QUYẾT TRANH CHẤP")
body("7.1. Hợp đồng chấm dứt khi thuộc một trong các trường hợp sau:")
bullet("Công việc đã hoàn thành và hai bên đã thanh toán, thanh lý đầy đủ;")
bullet("Hai bên thỏa thuận chấm dứt bằng văn bản;")
bullet("Một bên vi phạm nghiêm trọng nghĩa vụ của Hợp đồng và không khắc phục trong thời hạn 15 ngày kể từ ngày nhận thông báo bằng văn bản của bên còn lại.")
body("7.2. Bên vi phạm có trách nhiệm bồi thường thiệt hại thực tế phát sinh cho bên còn lại theo quy định của pháp luật.")
body("7.3. Mọi tranh chấp phát sinh từ hoặc liên quan đến Hợp đồng này được hai bên ưu tiên giải quyết bằng thương lượng, hòa giải. Trường hợp không thương lượng được, tranh chấp được đưa ra Tòa án nhân dân có thẩm quyền tại {dia_diem_giai_quyet_tranh_chap} để giải quyết theo quy định của pháp luật Việt Nam.")
add("")

# ---------- Dieu 8 ----------
heading("ĐIỀU 8. ĐIỀU KHOẢN CHUNG")
body("8.1. Hợp đồng này có hiệu lực kể từ ngày ký.")
body("8.2. Hai bên cam kết thực hiện đúng và đầy đủ các điều khoản đã thỏa thuận. Trong quá trình thực hiện, nếu có vướng mắc, hai bên kịp thời bàn bạc, thống nhất giải quyết trên tinh thần hợp tác.")
body("8.3. Mọi sửa đổi, bổ sung Hợp đồng chỉ có giá trị khi được lập thành văn bản và có chữ ký của đại diện có thẩm quyền của hai bên.")
body("8.4. Hợp đồng này được lập thành 04 bản có giá trị pháp lý như nhau; Bên A giữ 03 bản, Bên B giữ 01 bản.")
add("")
add("")

# ---------- Chu ky ----------
table = d.add_table(rows=1, cols=2)
table.autofit = True
hdr = table.rows[0].cells
for i, t in enumerate(["ĐẠI DIỆN BÊN A", "ĐẠI DIỆN BÊN B"]):
    hdr[i].text = ""
    p = hdr[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(13)
row = table.add_row().cells
for i, t in enumerate(["(Ký, ghi rõ họ tên, đóng dấu)", "(Ký, ghi rõ họ tên)"]):
    row[i].text = ""
    p = row[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    r.italic = True
    r.font.size = Pt(12)
for _ in range(4):
    row = table.add_row().cells
    for c in row:
        c.text = ""
row = table.add_row().cells
for i, t in enumerate(["{ben_a_dai_dien}", "{ben_b_ten}"]):
    row[i].text = ""
    p = row[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(13)

d.save(OUT)
print("Da ghi:", OUT)

# Liet ke placeholder de doi chieu
import re, zipfile
z = zipfile.ZipFile(OUT)
xml = z.read("word/document.xml").decode("utf-8")
tags = sorted(set(re.findall(r"\{([a-z0-9_]+)\}", xml)))
print("So placeholder:", len(tags))
print(", ".join(tags))
